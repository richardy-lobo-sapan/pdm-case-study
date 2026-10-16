import json
import re

import pytest
from pptx import Presentation

from src.data import ROOT
from src.field_note import numbers_in

DECK = ROOT / "deck" / "pdm_case_study.pptx"
FACTS = ROOT / "deck" / "facts.json"
pytestmark = pytest.mark.skipif(not (DECK.exists() and FACTS.exists()), reason="deck not built: run python deck/build_deck.py")


def _facts() -> dict:
    return json.loads(FACTS.read_text(encoding="utf-8"))


def _allowed() -> set[float]:
    """Every number that appears in a registered fact (each fact records where it came from)."""
    return {n for f in _facts().values() for n in numbers_in(f["text"])}


def _deck_text() -> str:
    prs = Presentation(DECK)
    parts = []
    for slide in prs.slides:
        parts += [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(parts)


def _prose(path) -> str:
    """Markdown without code blocks or inline code (commands, file names and versions are not results)."""
    text = path.read_text(encoding="utf-8")
    return re.sub(r"`[^`]*`", "", re.sub(r"```.*?```", "", text, flags=re.S))


@pytest.mark.parametrize("name, text", [
    ("deck", lambda: _deck_text()),
    ("LinkedIn post", lambda: _prose(ROOT / "deck" / "linkedin_post.md")),
    ("README", lambda: _prose(ROOT / "README.md")),
])
def test_every_number_shown_is_a_registered_fact(name, text):
    unknown = sorted(set(numbers_in(text())) - _allowed())
    assert not unknown, f"{name} shows numbers not produced by code: {unknown}"


def test_deck_is_not_stale():
    """The headline in the deck must match the current results.json (rebuild the deck after retraining)."""
    results = json.loads((ROOT / "reports" / "results.json").read_text(encoding="utf-8"))
    savings = f"{results['cost']['savings_lower_bound_policy']['savings_vs_best_alternative'] * 100:.0f}"
    assert _facts()["savings_pct"]["text"] == savings
    first_slide = Presentation(DECK).slides[0]
    assert any(f"{savings}%" in sh.text_frame.text for sh in first_slide.shapes if sh.has_text_frame)


def test_every_fact_names_its_source():
    assert all(f["source"].strip() for f in _facts().values())
